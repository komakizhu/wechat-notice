import tempfile
import unittest
from pathlib import Path
from importlib.util import find_spec

from wechat_notice.documents import extract_document


class DocumentTests(unittest.TestCase):
    def test_text_extraction_retains_a_source_location(self) -> None:
        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / "通知.txt"
            path.write_text("请在周五前提交名单。", encoding="utf-8")

            result = extract_document(path)

        self.assertEqual(result.status, "extracted")
        self.assertEqual(result.text, "[document] 请在周五前提交名单。")
        self.assertEqual(result.blocks[0].location, "document")

    def test_unsupported_file_is_explicit(self) -> None:
        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / "通知.zip"
            path.write_bytes(b"not parsed")
            result = extract_document(path)

        self.assertEqual(result.status, "unsupported")
        self.assertIn("zip", result.error or "")

    @unittest.skipUnless(find_spec("docx"), "python-docx not installed")
    def test_docx_extracts_paragraphs_and_tables(self) -> None:
        from docx import Document

        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / "通知.docx"
            document = Document()
            document.add_paragraph("请各班提交名单")
            table = document.add_table(rows=1, cols=2)
            table.cell(0, 0).text = "截止日期"
            table.cell(0, 1).text = "2026-09-20"
            document.save(path)

            result = extract_document(path)

        self.assertEqual(result.status, "extracted")
        self.assertIn("请各班提交名单", result.text)
        self.assertIn("table 1, row 1", result.text)

    @unittest.skipUnless(find_spec("openpyxl"), "openpyxl not installed")
    def test_xlsx_extracts_sheet_and_row_locations(self) -> None:
        from openpyxl import Workbook

        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / "名单.xlsx"
            workbook = Workbook()
            sheet = workbook.active
            sheet.title = "提交表"
            sheet.append(["姓名", "材料"])
            sheet.append(["张三", "已提交"])
            workbook.save(path)

            result = extract_document(path)

        self.assertEqual(result.status, "extracted")
        self.assertIn("sheet 提交表, row 2", result.text)
        self.assertIn("张三", result.text)

    @unittest.skipUnless(find_spec("pptx"), "python-pptx not installed")
    def test_pptx_extracts_slide_locations(self) -> None:
        from pptx import Presentation
        from pptx.util import Inches

        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / "安排.pptx"
            presentation = Presentation()
            slide = presentation.slides.add_slide(presentation.slide_layouts[6])
            box = slide.shapes.add_textbox(Inches(1), Inches(1), Inches(5), Inches(1))
            box.text = "活动安排"
            presentation.save(path)

            result = extract_document(path)

        self.assertEqual(result.status, "extracted")
        self.assertIn("slide 1, shape 1", result.text)
        self.assertIn("活动安排", result.text)

    @unittest.skipUnless(find_spec("pypdf") and find_spec("reportlab"), "PDF dependencies not installed")
    def test_pdf_extracts_page_locations(self) -> None:
        from reportlab.pdfgen import canvas

        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / "通知.pdf"
            writer = canvas.Canvas(str(path))
            writer.drawString(72, 720, "Submit before deadline")
            writer.save()

            result = extract_document(path)

        self.assertEqual(result.status, "extracted")
        self.assertIn("page 1", result.text)
        self.assertIn("Submit before deadline", result.text)

    def test_image_uses_injected_local_ocr_and_keeps_image_location(self) -> None:
        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / "公众号图片.png"
            path.write_bytes(b"not-a-real-image-for-injected-ocr")

            result = extract_document(path, ocr_runner=lambda _path: "图片里的报名截止时间")

        self.assertEqual(result.status, "extracted")
        self.assertEqual(result.blocks[0].location, "image")
        self.assertIn("图片里的报名截止时间", result.text)

    @unittest.skipUnless(find_spec("pypdf") and find_spec("reportlab"), "PDF dependencies not installed")
    def test_mixed_pdf_ocr_is_limited_to_pages_without_extractable_text(self) -> None:
        from reportlab.pdfgen import canvas

        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / "混合通知.pdf"
            writer = canvas.Canvas(str(path))
            writer.drawString(72, 720, "Text page")
            writer.showPage()
            writer.showPage()
            writer.save()

            rendered_pages: list[int] = []

            def renderer(_path: Path, page_number: int) -> bytes:
                rendered_pages.append(page_number)
                return b"rendered-page"

            result = extract_document(
                path,
                pdf_page_renderer=renderer,
                ocr_runner=lambda _path: "Scan page text",
            )

        self.assertEqual(result.status, "extracted")
        self.assertEqual(rendered_pages, [2])
        self.assertIn("page 1", result.text)
        self.assertIn("page 2 (OCR)", result.text)
        self.assertIn("Scan page text", result.text)


if __name__ == "__main__":
    unittest.main()
