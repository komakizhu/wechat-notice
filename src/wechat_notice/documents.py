from __future__ import annotations

import csv
import io
import json
import mimetypes
import tempfile
from dataclasses import dataclass
from pathlib import Path
from typing import Callable, Iterable

from .ocr import OCRResult, ocr_image_path


@dataclass(frozen=True)
class TextBlock:
    location: str
    text: str


@dataclass(frozen=True)
class DocumentExtraction:
    path: str
    name: str
    mime_type: str | None
    status: str
    blocks: tuple[TextBlock, ...] = ()
    error: str | None = None

    @property
    def text(self) -> str:
        return "\n".join(
            f"[{block.location}] {block.text}" for block in self.blocks if block.text.strip()
        )


def _result(path: Path, status: str, blocks: Iterable[TextBlock] = (), error: str | None = None) -> DocumentExtraction:
    return DocumentExtraction(
        path=str(path),
        name=path.name,
        mime_type=mimetypes.guess_type(path.name)[0],
        status=status,
        blocks=tuple(blocks),
        error=error,
    )


def _extract_text(path: Path) -> DocumentExtraction:
    try:
        if path.suffix.lower() == ".csv":
            with path.open("r", encoding="utf-8-sig", newline="") as handle:
                rows = csv.reader(handle)
                blocks = [TextBlock(f"row {index}", " | ".join(row)) for index, row in enumerate(rows, 1)]
        elif path.suffix.lower() == ".json":
            value = json.loads(path.read_text(encoding="utf-8"))
            blocks = [TextBlock("document", json.dumps(value, ensure_ascii=False, indent=2))]
        else:
            blocks = [TextBlock("document", path.read_text(encoding="utf-8"))]
    except (OSError, UnicodeError, json.JSONDecodeError) as exc:
        return _result(path, "failed", error=str(exc))
    return _result(path, "extracted", blocks)


def _render_pdf_page(path: Path, page_number: int) -> bytes:
    """Render one PDF page for OCR, using the optional local PDFium package."""
    import pypdfium2 as pdfium

    document = pdfium.PdfDocument(str(path))
    try:
        page = document[page_number - 1]
        bitmap = page.render(scale=2)
        image = bitmap.to_pil()
        output = io.BytesIO()
        image.save(output, format="PNG")
        return output.getvalue()
    finally:
        close = getattr(document, "close", None)
        if callable(close):
            close()


def _ocr_path_result(path: Path, runner: Callable[[Path], str | OCRResult]) -> OCRResult:
    try:
        result = runner(path)
    except Exception as exc:
        return OCRResult("failed", error=str(exc))
    if isinstance(result, OCRResult):
        return result
    text = str(result).strip()
    return OCRResult("extracted" if text else "empty", text=text)


def _ocr_page_bytes(
    data: bytes,
    page_number: int,
    runner: Callable[[Path], str | OCRResult],
) -> OCRResult:
    try:
        with tempfile.NamedTemporaryFile(
            suffix=f"-page-{page_number}.png", prefix="wechat-notice-"
        ) as handle:
            handle.write(data)
            handle.flush()
            return _ocr_path_result(Path(handle.name), runner)
    except OSError as exc:
        return OCRResult("failed", error=f"暂存 PDF 页面失败: {exc}")


def _extract_pdf(
    path: Path,
    *,
    ocr_runner: Callable[[Path], str | OCRResult] | None = None,
    pdf_page_renderer: Callable[[Path, int], bytes] | None = None,
) -> DocumentExtraction:
    try:
        from pypdf import PdfReader

        reader = PdfReader(str(path))
        blocks = []
        unresolved_pages: list[str] = []
        ocr = ocr_runner or ocr_image_path
        renderer = pdf_page_renderer or _render_pdf_page
        for index, page in enumerate(reader.pages, 1):
            text = page.extract_text() or ""
            if text.strip():
                blocks.append(TextBlock(f"page {index}", text.strip()))
                continue
            try:
                rendered = renderer(path, index)
                result = _ocr_page_bytes(rendered, index, ocr)
            except Exception as exc:
                result = OCRResult("failed", error=str(exc))
            if result.text.strip():
                blocks.append(TextBlock(f"page {index} (OCR)", result.text.strip()))
            else:
                unresolved_pages.append(
                    f"第{index}页（{result.error or '未识别出文字'}）"
                )
        if unresolved_pages and not blocks:
            return _result(
                path,
                "needs_visual_review",
                error="PDF 扫描页未完成 OCR：" + "、".join(unresolved_pages),
            )
        combined = "\n".join(block.text for block in blocks)
        replacement_count = sum(combined.count(marker) for marker in ("\ufffd", "■", "□"))
        if replacement_count and replacement_count / max(len(combined), 1) >= 0.2:
            return _result(
                path,
                "needs_review",
                blocks,
                error="PDF 文本可能因字体编码产生乱码，需要人工核对或提供可读版本",
            )
        if unresolved_pages:
            return _result(
                path,
                "needs_visual_review",
                blocks,
                error="PDF 部分页面未完成 OCR：" + "、".join(unresolved_pages),
            )
        return _result(path, "extracted", blocks)
    except Exception as exc:  # parser errors vary across pypdf versions
        return _result(path, "failed", error=f"PDF 解析失败: {exc}")


def _extract_docx(path: Path) -> DocumentExtraction:
    try:
        from docx import Document

        document = Document(str(path))
        blocks: list[TextBlock] = []
        for index, paragraph in enumerate(document.paragraphs, 1):
            text = paragraph.text.strip()
            if text:
                blocks.append(TextBlock(f"paragraph {index}", text))
        for table_index, table in enumerate(document.tables, 1):
            for row_index, row in enumerate(table.rows, 1):
                values = [cell.text.strip() for cell in row.cells]
                if any(values):
                    blocks.append(TextBlock(f"table {table_index}, row {row_index}", " | ".join(values)))
        if not blocks:
            return _result(path, "no_text", error="Word 文件没有可提取文本")
        return _result(path, "extracted", blocks)
    except Exception as exc:
        return _result(path, "failed", error=f"DOCX 解析失败: {exc}")


def _extract_xlsx(path: Path) -> DocumentExtraction:
    try:
        from openpyxl import load_workbook

        workbook = load_workbook(str(path), read_only=True, data_only=False)
        blocks: list[TextBlock] = []
        for sheet in workbook.worksheets:
            for row_index, row in enumerate(sheet.iter_rows(), 1):
                values = ["" if cell.value is None else str(cell.value) for cell in row]
                if any(value.strip() for value in values):
                    blocks.append(TextBlock(f"sheet {sheet.title}, row {row_index}", " | ".join(values)))
        workbook.close()
        if not blocks:
            return _result(path, "no_text", error="Excel 文件没有可提取内容")
        return _result(path, "extracted", blocks)
    except Exception as exc:
        return _result(path, "failed", error=f"XLSX 解析失败: {exc}")


def _extract_pptx(path: Path) -> DocumentExtraction:
    try:
        from pptx import Presentation

        presentation = Presentation(str(path))
        blocks: list[TextBlock] = []
        for slide_index, slide in enumerate(presentation.slides, 1):
            for shape_index, shape in enumerate(slide.shapes, 1):
                text = getattr(shape, "text", "").strip()
                if text:
                    blocks.append(TextBlock(f"slide {slide_index}, shape {shape_index}", text))
        if not blocks:
            return _result(path, "no_text", error="PowerPoint 文件没有可提取文本")
        return _result(path, "extracted", blocks)
    except Exception as exc:
        return _result(path, "failed", error=f"PPTX 解析失败: {exc}")


def _extract_image(
    path: Path,
    *,
    ocr_runner: Callable[[Path], str | OCRResult] | None = None,
) -> DocumentExtraction:
    result = _ocr_path_result(path, ocr_runner or ocr_image_path)
    if result.text.strip():
        return _result(path, "extracted", [TextBlock("image", result.text.strip())])
    status = "needs_visual_review" if result.status in {"unavailable", "empty"} else "failed"
    return _result(
        path,
        status,
        error=result.error or "图片未识别出文字；需要视觉核对",
    )


def extract_document(
    path: str | Path,
    *,
    ocr_runner: Callable[[Path], str | OCRResult] | None = None,
    pdf_page_renderer: Callable[[Path, int], bytes] | None = None,
) -> DocumentExtraction:
    """Extract text with source locations without modifying the source file."""
    candidate = Path(path).expanduser()
    if not candidate.is_file():
        return _result(candidate, "failed", error="文件不存在或不是普通文件")

    suffix = candidate.suffix.lower()
    if suffix in {".txt", ".md", ".markdown", ".csv", ".json"}:
        return _extract_text(candidate)
    if suffix == ".pdf":
        return _extract_pdf(
            candidate,
            ocr_runner=ocr_runner,
            pdf_page_renderer=pdf_page_renderer,
        )
    if suffix == ".docx":
        return _extract_docx(candidate)
    if suffix in {".xlsx", ".xlsm"}:
        return _extract_xlsx(candidate)
    if suffix == ".pptx":
        return _extract_pptx(candidate)
    if suffix in {".png", ".jpg", ".jpeg", ".webp", ".gif", ".bmp", ".tif", ".tiff"}:
        return _extract_image(candidate, ocr_runner=ocr_runner)
    return _result(candidate, "unsupported", error=f"暂不支持的文件类型: {suffix or '无扩展名'}")
